"""
reporting.py — Excel exception pack for the whole engagement, quarter by quarter.
"""
import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from models import (Txn, Match, Split, Exception_, DocRequest, DocumentFile,
                    Balance, Task, QuarterFile, Account)
import core

NAVY = PatternFill("solid", fgColor="1F3864")
RED = PatternFill("solid", fgColor="F4CCCC")
HF = Font(bold=True, color="FFFFFF", size=10)
BF = Font(size=10)
TF = Font(bold=True, size=13, color="1F3864")


def _sheet(wb, title, headers, rows, red_col=None, red_val=None):
    ws = wb.create_sheet(title[:31])
    for j, h in enumerate(headers, 1):
        c = ws.cell(1, j, h)
        c.fill, c.font = NAVY, HF
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    for i, row in enumerate(rows, 2):
        flag = red_col is not None and str(row[red_col]) == red_val
        for j, v in enumerate(row, 1):
            c = ws.cell(i, j, v)
            c.font = BF
            if flag:
                c.fill = RED
    for col in ws.columns:
        w = max((len(str(c.value)) for c in col if c.value is not None), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max(w + 2, 10), 38)
    ws.freeze_panes = "A2"
    return ws


def build_workbook(engagement) -> bytes:
    eid = engagement.id
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = f"Verity — {engagement.name}"
    ws["A1"].font = TF
    ws["A2"] = f"{engagement.client or ''} | {engagement.currency} | {engagement.period_start} to {engagement.period_end}"
    ws["A2"].font = BF

    qs = core.quarters(eid)
    hdr = ("Quarter", "Months", "Status", "Matched", "Exceptions", "Open", "Value open", "Signed off")
    for j, h in enumerate(hdr, 1):
        c = ws.cell(4, j, h)
        c.fill, c.font = NAVY, HF
    qfs = {q.quarter: q for q in QuarterFile.query.filter_by(engagement_id=eid).all()}
    for i, q in enumerate(qs, 5):
        exc = Exception_.query.filter_by(engagement_id=eid, quarter=q["quarter"]).all()
        vals = (q["quarter"], len(q["months"]), q["status"],
                Match.query.filter_by(engagement_id=eid, quarter=q["quarter"]).count(),
                len(exc), sum(1 for x in exc if x.status == "Open"),
                round(sum(x.amount or 0 for x in exc if x.status != "Cleared"), 0),
                (qfs.get(q["quarter"]).signed_off_by if qfs.get(q["quarter"]) else "") or "")
        for j, v in enumerate(vals, 1):
            ws.cell(i, j, v).font = BF
    for col, w in zip("ABCDEFGH", [12, 9, 13, 10, 12, 8, 14, 18]):
        ws.column_dimensions[col].width = w

    _sheet(wb, "Exceptions",
           ["Ref", "Quarter", "Category", "Date", "Counterparty", "Amount", "Status", "Detail", "Resolution"],
           [(x.ref, x.quarter, x.category, x.date, x.counterparty, x.amount, x.status,
             x.detail, x.resolution) for x in
            Exception_.query.filter_by(engagement_id=eid).order_by(Exception_.quarter).all()],
           red_col=6, red_val="Open")

    _sheet(wb, "Document_requests",
           ["Ref", "Quarter", "Addressee", "Documents", "Raised", "Due", "Status", "Chases", "Response"],
           [(r.ref, r.quarter, r.addressee, r.documents_needed,
             r.raised_at.date() if r.raised_at else None, r.due_date, r.status,
             r.chased_count, r.response_note) for r in
            DocRequest.query.filter_by(engagement_id=eid).all()])

    _sheet(wb, "Balance_continuity",
           ["Account", "Month", "Opening", "Closing", "Movement", "Variance", "Continuity gap", "Anomaly"],
           [(b.account.code if b.account else "", b.month, b.opening, b.closing,
             b.computed_movement, b.variance, b.continuity_gap, b.anomaly) for b in
            Balance.query.filter_by(engagement_id=eid).order_by(Balance.month).all()])

    _sheet(wb, "Match_quality",
           ["Quarter", "Month", "Tier", "How matched", "Confidence", "Amount", "Date gap"],
           [(m.quarter, m.month, m.tier, m.tier_label, m.confidence, m.amount, m.date_gap_days)
            for m in Match.query.filter_by(engagement_id=eid).order_by(Match.quarter).all()])

    _sheet(wb, "Split_posting", ["Quarter", "Pattern", "Amount", "Counterparty", "Linked"],
           [(s.quarter, s.pattern, s.amount, s.counterparty, s.many_txn_ids)
            for s in Split.query.filter_by(engagement_id=eid).all()])

    _sheet(wb, "Documents", ["Ref", "Filename", "Type", "Quarter", "Txn reference", "Uploaded by"],
           [(d.ref, d.filename, d.doc_type, d.quarter, d.txn_reference, d.uploaded_by)
            for d in DocumentFile.query.filter_by(engagement_id=eid).all()])

    _sheet(wb, "Quarter_rationale", ["Quarter", "Value at risk", "Signed off by", "Narrative"],
           [(q.quarter, q.value_at_risk, q.signed_off_by, q.narrative)
            for q in QuarterFile.query.filter_by(engagement_id=eid).order_by(QuarterFile.quarter).all()])

    _sheet(wb, "Tasks", ["Ref", "Phase", "Task", "Owner", "Joint", "Start", "Due", "Status", "%"],
           [(t.ref, t.phase, t.title, t.owner.name if t.owner else "",
             t.joint_owner.name if t.joint_owner else "", t.start_date, t.due_date,
             t.status, t.progress) for t in
            Task.query.filter_by(engagement_id=eid).order_by(Task.due_date).all()])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()
